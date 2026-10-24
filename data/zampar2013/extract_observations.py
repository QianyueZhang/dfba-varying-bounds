"""Extract published observations, preserving workbook values and provenance.

No regression, model prediction, digitisation or imputation is performed.
Run with a Python environment containing openpyxl and matplotlib.
"""
from pathlib import Path
from collections import Counter, defaultdict
import csv
import hashlib
import json
import math
import re
import zipfile

import openpyxl

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "supplementary_file_1" / "Regressions"
VARIABLES = {
    "Biomass": ("biomass_OD600", "OD600 (AU)", "derived_dry_weight", "gDW/L"),
    "Glucose": ("glucose", "g/L", "converted_concentration", "mM"),
    "Ethanol": ("ethanol", "g/L", "converted_concentration", "mM"),
}
archive = ROOT / "msb201311-s1.zip"
assert hashlib.sha256(archive.read_bytes()).hexdigest() == "d928c613be23fe67e6af1ec2e75c88390a4a48b7824c349fca736081c5cecaeb"
if not all((SOURCE / f"{name}.xlsx").exists() for name in VARIABLES):
    with zipfile.ZipFile(archive) as source_zip:
        for name in VARIABLES:
            source_zip.extract(f"Regressions/{name}.xlsx", ROOT / "supplementary_file_1")
rows = []
summary = {}
for filename, (variable, unit, secondary_kind, secondary_unit) in VARIABLES.items():
    path = SOURCE / f"{filename}.xlsx"
    worksheet = openpyxl.load_workbook(path, data_only=True)["Overview"]
    assert worksheet["C10"].value == "Experimental"
    label = None
    repeats = Counter()
    for row_num in range(13, worksheet.max_row + 1):
        possible_label = worksheet.cell(row_num, 2).value
        if isinstance(possible_label, str) and re.fullmatch(r"Exp\s*\d+", possible_label):
            label = possible_label
        time, value, secondary = [worksheet.cell(row_num, c).value for c in (3, 4, 5)]
        if not isinstance(time, (int, float)):
            continue
        assert label is not None
        # The source's experimental block is C:E. F:G and I:L are excluded.
        assert isinstance(value, (int, float)), (filename, row_num, value)
        assert isinstance(secondary, (int, float)), (filename, row_num, secondary)
        key = (label, time)
        repeats[key] += 1
        rows.append({
            "source_doi": "10.1038/msb.2013.11",
            "source_file": f"Regressions/{filename}.xlsx",
            "source_sheet": "Overview",
            "source_row": row_num,
            "source_time_cell": f"C{row_num}",
            "source_value_cell": f"D{row_num}",
            "source_secondary_cell": f"E{row_num}",
            "variable": variable,
            "replicate_label_as_published": label,
            "series_id": f"{filename}_" + label.replace(" ", ""),
            "cross_variable_replicate_identity_verified": False,
            "time_relative_to_glucose_exhaustion_h": time,
            "repeat_index_at_same_series_time": repeats[key],
            "observed_value_as_published": value,
            "observed_unit": unit,
            "secondary_value_as_published": secondary,
            "secondary_value_kind": secondary_kind,
            "secondary_unit": secondary_unit,
            "missing": False,
            "reported_zero": value == 0,
            "lod_value": "",
            "lod_status": "not_reported",
            "censoring_status": "unspecified_for_reported_zero" if value == 0 else "not_flagged_by_source",
        })
    own = [r for r in rows if r["variable"] == variable]
    by_series = defaultdict(list)
    for r in own:
        by_series[r["series_id"]].append(r)
    summary[variable] = {
        "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "n_observations": len(own),
        "n_reported_zeros": sum(r["reported_zero"] for r in own),
        "n_missing_in_selected_experimental_rows": 0,
        "range_h": [min(r["time_relative_to_glucose_exhaustion_h"] for r in own), max(r["time_relative_to_glucose_exhaustion_h"] for r in own)],
        "unit": unit,
        "series": [{
            "series_id": key, "n": len(group),
            "earliest_time_h": min(r["time_relative_to_glucose_exhaustion_h"] for r in group),
            "latest_time_h": max(r["time_relative_to_glucose_exhaustion_h"] for r in group),
            "earliest_observations": [{"time_h": r["time_relative_to_glucose_exhaustion_h"], "value": r["observed_value_as_published"], "source_cell": r["source_value_cell"]} for r in group if r["time_relative_to_glucose_exhaustion_h"] == min(s["time_relative_to_glucose_exhaustion_h"] for s in group)],
        } for key, group in by_series.items()],
    }

# Check source dry weights without changing either published representation.
biomass = [r for r in rows if r["variable"] == "biomass_OD600"]
assert all(math.isclose(r["secondary_value_as_published"], 0.485*r["observed_value_as_published"], abs_tol=1e-12) for r in biomass)
summary["biomass_conversion"] = {
    "paper_methods_gDW_per_L_per_OD600": 0.486,
    "factor_matching_all_workbook_experimental_dry_weights": 0.485,
    "relative_difference_percent": (0.486/0.485-1)*100,
    "decision": "Preserve raw OD and workbook derived dry weights separately. Plot raw OD; do not silently replace source dry weights.",
}
with (ROOT / "zampar2013_observations.csv").open("w", newline="") as handle:
    writer = csv.DictWriter(handle, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
(ROOT / "observation_summary.json").write_text(json.dumps(summary, indent=2))

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False})
fig, axes = plt.subplots(3, 1, figsize=(8.2, 8.6), sharex=True)
colours = ["#0072B2", "#D55E00", "#009E73", "#CC79A7"]
markers = ["o", "s", "^", "D"]
for ax, (filename, (variable, unit, _, _)) in zip(axes, VARIABLES.items()):
    own = [r for r in rows if r["variable"] == variable]
    labels = list(dict.fromkeys(r["replicate_label_as_published"] for r in own))
    for i, label in enumerate(labels):
        group = [r for r in own if r["replicate_label_as_published"] == label]
        ax.scatter([r["time_relative_to_glucose_exhaustion_h"] for r in group], [r["observed_value_as_published"] for r in group], color=colours[i], marker=markers[i], s=24, linewidths=0.5, edgecolors="white", alpha=0.85, label=label)
    ax.axvline(0, color="#888888", linestyle="--", linewidth=0.9)
    ax.grid(alpha=0.16)
    ax.set_ylabel("Biomass (OD$_{600}$, AU)" if filename == "Biomass" else f"{filename} (g L$^{{-1}}$)")
    ax.legend(title="Workbook series", fontsize=8, title_fontsize=8, ncol=len(labels), loc="lower left", bbox_to_anchor=(0, 1.01), borderaxespad=0, frameon=False)
    ax.set_ylim(bottom=-0.04 * max(r["observed_value_as_published"] for r in own))
axes[-1].set_xlabel("Time relative to glucose exhaustion (h)")
fig.suptitle("Published observations: Zampar et al. (2013)", y=0.98, fontsize=13)
fig.text(0.08, 0.022, "Experimental columns only; no regression or thesis predictions.\nSeries labels are workbook-specific; cross-variable replicate matching has not been established.", fontsize=8, color="#444444")
fig.tight_layout(rect=(0, 0.07, 1, 0.96))
fig.savefig(ROOT / "zampar2013_observations.png", dpi=220)
fig.savefig(ROOT / "zampar2013_observations.svg")
print(json.dumps({"total_observations":len(rows), "counts":{key:val["n_observations"] for key,val in summary.items() if "n_observations" in val}, "output":str(ROOT)}, indent=2))
